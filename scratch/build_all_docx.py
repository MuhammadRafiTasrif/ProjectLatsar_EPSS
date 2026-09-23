import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import os

def create_styled_document():
    doc = docx.Document()
    
    # Page Setup - Margins 1 inch (2.54 cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    return doc

def add_header_banner(doc, title_text, subtitle_text):
    # Header Banner Block
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_header.add_run("BADAN PUSAT STATISTIK KABUPATEN PASAMAN\nPROVINSI SUMATERA BARAT\n")
    run_inst.font.name = 'Calibri'
    run_inst.font.size = Pt(11)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(30, 41, 59)

    run_line = p_header.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━\n")
    run_line.font.color.rgb = RGBColor(37, 99, 235) # Primary Blue
    run_line.font.bold = True

    run_title = p_header.add_run(f"{title_text}\n")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(29, 78, 216) # Dark Blue

    run_sub = p_header.add_run(f"{subtitle_text}\n")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph() # Spacing

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Dark Navy
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(37, 99, 235) # Royal Blue
    return p

def add_callout(doc, text, title="📌 CATATAN PENTING"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    # Set shading light blue background
    shading_elm = parse_xml(r'<w:shd {} w:fill="EFF6FF"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading_elm)

    # Set borders: left thick blue border, others none
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="2563EB"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    '''.format(nsdecls('w')))
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"{title}\n")
    r_title.font.bold = True
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor(29, 78, 216)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph() # Spacing

def format_table_headers_and_borders(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(r'<w:tblHeader {} />'.format(nsdecls('w'))))

    for cell in table.rows[0].cells:
        shading = parse_xml(r'<w:shd {} w:fill="1E3A8A"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    for row_idx, row in enumerate(table.rows[1:], start=1):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit {} />'.format(nsdecls('w'))))
        
        # Zebra striping for even rows
        if row_idx % 2 == 0:
            for cell in row.cells:
                shd = parse_xml(r'<w:shd {} w:fill="F8FAFC"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(shd)

        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(30, 41, 59)

    # Table borders
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(r'''
        <w:tblBorders {} >
            <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="94A3B8"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    '''.format(nsdecls('w')))
    tblPr.append(tblBorders)

print("Docx generator library helper initialized.")
