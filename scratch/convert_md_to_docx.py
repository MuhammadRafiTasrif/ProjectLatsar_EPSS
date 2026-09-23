import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

MD_PATH = r"d:\3. BPS\1. 1309_(D) Laptop\1. Project Latsar\Aplikasi 1\laporan\Buku_Pedoman_Penggunaan_Sistem_SIMPONITAS.md"
DOCX_PATH = r"d:\3. BPS\1. 1309_(D) Laptop\1. Project Latsar\Aplikasi 1\laporan\Buku_Pedoman_Penggunaan_Sistem_SIMPONITAS.docx"

doc = docx.Document()

# Set standard margins (2.5 cm all around)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Cover Header
p_inst = doc.add_paragraph()
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_inst = p_inst.add_run("BADAN PUSAT STATISTIK KABUPATEN PASAMAN\nPROVINSI SUMATERA BARAT\n")
run_inst.font.name = "Calibri"
run_inst.font.size = Pt(13)
run_inst.font.bold = True
run_inst.font.color.rgb = RGBColor(15, 23, 42)

p_rule = doc.add_paragraph()
p_rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_rule_run = p_rule.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
p_rule_run.font.color.rgb = RGBColor(247, 144, 57)

# Title
p_space = doc.add_paragraph("\n\n")
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_badge = p_title.add_run("BUKU PEDOMAN RESMI PENGGUNA\n\n")
run_badge.font.name = "Calibri"
run_badge.font.size = Pt(12)
run_badge.font.bold = True
run_badge.font.color.rgb = RGBColor(194, 65, 12)

run_main = p_title.add_run("PANDUAN PENGGUNAAN SISTEM INFORMASI SIMPONITAS\n")
run_main.font.name = "Calibri"
run_main.font.size = Pt(20)
run_main.font.bold = True
run_main.font.color.rgb = RGBColor(15, 23, 42)

run_sub = p_title.add_run('"Sinergi Pembinaan Statistik Sektoral melalui Penerbitan Kompromin Akurat dan Berkualitas"\n\n\n\n')
run_sub.font.name = "Calibri"
run_sub.font.size = Pt(12)
run_sub.font.italic = True
run_sub.font.color.rgb = RGBColor(247, 144, 57)

# Author
p_author = doc.add_paragraph()
p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_author = p_author.add_run("Disusun Oleh:\n")
run_author.font.name = "Calibri"
run_author.font.size = Pt(11)
run_author_name = p_author.add_run("MUHAMMAD RAFI TASRIF, S.Tr.Stat\n")
run_author_name.font.name = "Calibri"
run_author_name.font.size = Pt(12)
run_author_name.font.bold = True
run_author_desc = p_author.add_run("NIP. 20010925 202404 1 001\nPranata Komputer Ahli Pertama — BPS Kabupaten Pasaman\n\nAktualisasi Pelatihan Dasar CPNS Golongan III BPS Tahun 2026")
run_author_desc.font.name = "Calibri"
run_author_desc.font.size = Pt(10)
run_author_desc.font.color.rgb = RGBColor(100, 116, 139)

doc.add_page_break()

# Read markdown content line by line and construct document
with open(MD_PATH, "r", encoding="utf-8") as f:
    lines = f.readlines()

in_table = False
table_rows = []

def flush_table(rows):
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=col_count)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx, cell_value in enumerate(row):
            if c_idx < len(tbl.rows[r_idx].cells):
                cell = tbl.rows[r_idx].cells[c_idx]
                cell.text = cell_value.strip()
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
                    if r_idx == 0:
                        run.font.bold = True
                if r_idx == 0:
                    shd = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
                    cell._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

for line in lines:
    raw = line.rstrip()
    
    # Table detection
    if raw.startswith("|") and raw.endswith("|"):
        # Check if separator row
        if set(raw.replace("|", "").strip()) <= {"-", ":", " "}:
            continue
        parts = [c.strip() for c in raw.strip("|").split("|")]
        table_rows.append(parts)
        in_table = True
        continue
    else:
        if in_table:
            flush_table(table_rows)
            table_rows = []
            in_table = False
            
    if not raw:
        continue
        
    if raw.startswith("# "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(raw[2:])
        r.font.name = "Calibri"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
    elif raw.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(raw[3:])
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
    elif raw.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(raw[4:])
        r.font.name = "Calibri"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(51, 65, 85)
    elif raw.startswith("#### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(raw[5:])
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(71, 85, 105)
    elif raw.startswith("* ") or raw.startswith("- "):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        text = raw[2:]
        # Simple bold formatting parse
        parts = text.split("**")
        for idx, part in enumerate(parts):
            r = p.add_run(part)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            if idx % 2 == 1:
                r.font.bold = True
    elif raw.startswith("1. ") or raw.startswith("2. ") or raw.startswith("3. ") or raw.startswith("4. ") or raw.startswith("5. ") or raw.startswith("6. ") or raw.startswith("7. ") or raw.startswith("8. ") or raw.startswith("9. "):
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        text = raw[3:]
        parts = text.split("**")
        for idx, part in enumerate(parts):
            r = p.add_run(part)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            if idx % 2 == 1:
                r.font.bold = True
    elif raw.startswith("```"):
        continue
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        parts = raw.split("**")
        for idx, part in enumerate(parts):
            r = p.add_run(part)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            if idx % 2 == 1:
                r.font.bold = True

if in_table:
    flush_table(table_rows)

doc.save(DOCX_PATH)
print("DOCX generated successfully at:", DOCX_PATH)
