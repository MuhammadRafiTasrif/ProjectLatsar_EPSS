import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

MD_PATH = r"d:\3. BPS\1. 1309_(D) Laptop\1. Project Latsar\Aplikasi 1\laporan\Buku_Pedoman_Penggunaan_Sistem_SIMPONITAS.md"
DOCX_PATH = r"d:\3. BPS\1. 1309_(D) Laptop\1. Project Latsar\Aplikasi 1\laporan\Buku_Pedoman_Penggunaan_Sistem_SIMPONITAS.docx"

doc = docx.Document()

# Set page margins to standard A4 (2.54 cm all around)
for section in doc.sections:
    section.page_width = Inches(8.27)  # A4 width
    section.page_height = Inches(11.69) # A4 height
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Configure Header & Footer
    section.different_first_page_header_footer = True
    header = section.header
    p_hdr = header.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hdr = p_hdr.add_run("Buku Pedoman Penggunaan Sistem SIMPONITAS — BPS Kabupaten Pasaman")
    r_hdr.font.name = "Calibri"
    r_hdr.font.size = Pt(8.5)
    r_hdr.font.italic = True
    r_hdr.font.color.rgb = RGBColor(148, 163, 184)
    
    footer = section.footer
    p_ftr = footer.paragraphs[0]
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ftr = p_ftr.add_run("Badan Pusat Statistik Kabupaten Pasaman — Aktualisasi Latsar CPNS 2026")
    r_ftr.font.name = "Calibri"
    r_ftr.font.size = Pt(8.5)
    r_ftr.font.color.rgb = RGBColor(148, 163, 184)

# =========================================================================
# COVER PAGE
# =========================================================================
p_inst = doc.add_paragraph()
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_inst.paragraph_format.space_before = Pt(20)
p_inst.paragraph_format.space_after = Pt(2)
r_inst = p_inst.add_run("BADAN PUSAT STATISTIK KABUPATEN PASAMAN\n")
r_inst.font.name = "Calibri"
r_inst.font.size = Pt(14)
r_inst.font.bold = True
r_inst.font.color.rgb = RGBColor(15, 23, 42)

r_subinst = p_inst.add_run("PROVINSI SUMATERA BARAT\n")
r_subinst.font.name = "Calibri"
r_subinst.font.size = Pt(11)
r_subinst.font.bold = True
r_subinst.font.color.rgb = RGBColor(100, 116, 139)

p_rule = doc.add_paragraph()
p_rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_rule.paragraph_format.space_after = Pt(40)
r_rule = p_rule.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
r_rule.font.color.rgb = RGBColor(247, 144, 57)

# Title Area
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_after = Pt(8)
r_badge = p_title.add_run("BUKU PEDOMAN RESMI PENGGUNA\n\n")
r_badge.font.name = "Calibri"
r_badge.font.size = Pt(12)
r_badge.font.bold = True
r_badge.font.color.rgb = RGBColor(194, 65, 12)

r_main = p_title.add_run("PANDUAN PENGGUNAAN SISTEM INFORMASI\nSIMPONITAS\n")
r_main.font.name = "Calibri"
r_main.font.size = Pt(22)
r_main.font.bold = True
r_main.font.color.rgb = RGBColor(15, 23, 42)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_before = Pt(8)
p_sub.paragraph_format.space_after = Pt(60)
r_sub = p_sub.add_run('"Sinergi Pembinaan Statistik Sektoral melalui Penerbitan Kompromin Akurat dan Berkualitas"')
r_sub.font.name = "Calibri"
r_sub.font.size = Pt(12)
r_sub.font.italic = True
r_sub.font.color.rgb = RGBColor(247, 144, 57)

# Author Area
p_author = doc.add_paragraph()
p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_author.paragraph_format.space_before = Pt(80)
p_author.paragraph_format.space_after = Pt(2)
r_auth_lbl = p_author.add_run("Disusun Oleh:\n")
r_auth_lbl.font.name = "Calibri"
r_auth_lbl.font.size = Pt(11)
r_auth_lbl.font.color.rgb = RGBColor(71, 85, 105)

r_auth_name = p_author.add_run("MUHAMMAD RAFI TASRIF, S.Tr.Stat\n")
r_auth_name.font.name = "Calibri"
r_auth_name.font.size = Pt(13)
r_auth_name.font.bold = True
r_auth_name.font.color.rgb = RGBColor(15, 23, 42)

r_auth_desc = p_author.add_run("NIP. 20010925 202404 1 001\nPranata Komputer Ahli Pertama\nBadan Pusat Statistik Kabupaten Pasaman\n\nAktualisasi Pelatihan Dasar CPNS Golongan III BPS Tahun 2026")
r_auth_desc.font.name = "Calibri"
r_auth_desc.font.size = Pt(10)
r_auth_desc.font.color.rgb = RGBColor(100, 116, 139)

doc.add_page_break()

# =========================================================================
# BODY PARSER
# =========================================================================
with open(MD_PATH, "r", encoding="utf-8") as f:
    lines = f.readlines()

in_table = False
table_rows = []

def flush_table(rows):
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=col_count)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Apply borders via oxml
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(r'''
        <w:tblBorders {} >
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    '''.format(nsdecls('w')))
    tblPr.append(borders)

    for r_idx, row in enumerate(rows):
        for c_idx, cell_value in enumerate(row):
            if c_idx < len(tbl.rows[r_idx].cells):
                cell = tbl.rows[r_idx].cells[c_idx]
                cell.text = cell_value.strip()
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
                    if r_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(15, 23, 42)
                    else:
                        run.font.color.rgb = RGBColor(51, 65, 85)
                        
                if r_idx == 0:
                    shd = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
                    cell._tc.get_or_add_tcPr().append(shd)
                    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

for line in lines:
    raw = line.rstrip()
    
    # Table detection
    if raw.startswith("|") and raw.endswith("|"):
        # Check if markdown separator row
        if set(raw.replace("|", "").strip()) <= {"-", ":", " "}:
            continue
        parts = [c.strip() for c in raw.strip("|").split("|")]
        table_rows.append(parts)
        in_table = True
        continue
    else:
        if in_table:
            flush_table(table_rows)
            table_rows = []
            in_table = False
            
    if not raw:
        continue
        
    # Headings
    if raw.startswith("# "):
        # Major Chapter
        heading_text = raw[2:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(heading_text)
        r.font.name = "Calibri"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        
        # Add decorative line under Chapter 1-10
        p_line = doc.add_paragraph()
        p_line.paragraph_format.space_before = Pt(0)
        p_line.paragraph_format.space_after = Pt(8)
        r_line = p_line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
        r_line.font.size = Pt(8)
        r_line.font.color.rgb = RGBColor(247, 144, 57)
        
    elif raw.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(raw[3:].strip())
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
    elif raw.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(raw[4:].strip())
        r.font.name = "Calibri"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(51, 65, 85)
    elif raw.startswith("#### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(raw[5:].strip())
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(71, 85, 105)
    elif raw.startswith("* ") or raw.startswith("- "):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        text = raw[2:].strip()
        parts = text.split("**")
        for idx, part in enumerate(parts):
            r = p.add_run(part)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            if idx % 2 == 1:
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
    elif len(raw) > 3 and raw[0].isdigit() and raw[1:3] in ['. ', ') ']:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        text = raw[3:].strip()
        parts = text.split("**")
        for idx, part in enumerate(parts):
            r = p.add_run(part)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            if idx % 2 == 1:
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
    elif raw.startswith("```"):
        continue
    elif raw.startswith("---"):
        continue
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        parts = raw.split("**")
        for idx, part in enumerate(parts):
            r = p.add_run(part)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            if idx % 2 == 1:
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)

if in_table:
    flush_table(table_rows)

doc.save(DOCX_PATH)
print("SUCCESS: Official DOCX created at:", DOCX_PATH)
